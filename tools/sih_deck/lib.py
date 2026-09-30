"""Small drawing layer over python-pptx for the SIH idea deck (inches, hex colours, rich text)."""
from __future__ import annotations

import json
import os
import re
import subprocess
from copy import deepcopy

from lxml import etree
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ICON_DIR = os.path.join(HERE, ".icons")
NODE_ICONS = os.path.join(HERE, "icons.cjs")

# ---------------------------------------------------------------- tokens (mirror frontend/app/globals.css)
INK, INK2, INK3, INK4 = "0C0E14", "3D4452", "555C6B", "6A7180"
RULE, RULE2 = "CDD2DC", "E3E6EC"
TINT, TINT2 = "F5F6F9", "ECEEF3"
WHITE = "FFFFFF"
BRAND, BRAND_INK, BRAND_SOFT = "5B4BFF", "4533EA", "EEEBFF"
AMBER, AMBER_DEEP, AMBER_SOFT = "F2A60C", "93530A", "FDF1D6"
CRIT, CRIT_DEEP, CRIT_SOFT = "D92D32", "B91C1F", "FCE8E8"
HIGH, HIGH_DEEP, HIGH_SOFT = "F76B15", "B93E0B", "FFF1E6"
LOW, LOW_DEEP, LOW_SOFT = "12A679", "0B7A58", "E4F7F0"
FAM = {"mud": "0EA5E9", "stuck": "D946EF", "kick": "EF4444", "torque": "EAB308", "instab": "14B8A6", "cement": "64748B"}
FORM = {"Girujan": "7B86A3", "Tipam": "E2A13B", "Namsang": "8FB34E", "Barail": "2F9FC2", "Kopili": "8D68D6", "Sylhet": "1FAE86"}
SIH_BLUE = "0070C0"

DISPLAY = "Franklin Gothic Demi Cond"  # heavy semi-condensed grotesk (static TTF, embeds in PDF)
DISPLAY_MED = "Franklin Gothic Medium Cond"
BODY = "Segoe UI"
BODY_SB = "Segoe UI Semibold"
BODY_B = "Segoe UI Bold"
MONO = "Consolas"

_icon_queue: list[dict] = []


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


def _strip_style(shape) -> None:
    """Drop the theme <p:style> so no theme shadow/colour leaks into our shapes."""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def _ln(shape):
    return shape._element.spPr.get_or_add_ln()


def set_line(shape, color: str | None, width: float = 0.75, dash: str | None = None) -> None:
    if color is None:
        shape.line.fill.background()
        return
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    if dash:
        ln = _ln(shape)
        for old in ln.findall(qn("a:prstDash")):
            ln.remove(old)
        pd = etree.SubElement(ln, qn("a:prstDash"))
        pd.set("val", dash)
        # prstDash must follow the fill element
        fill = ln.find(qn("a:solidFill"))
        if fill is not None:
            fill.addnext(pd)


def hard_shadow(shape, dist_pt: float = 3.0, color: str = INK) -> None:
    """Flat offset shadow (no blur) toward bottom-right — the deck's one decorative device."""
    spPr = shape._element.spPr
    for old in spPr.findall(qn("a:effectLst")):
        spPr.remove(old)
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"))
    sh.set("blurRad", "0")
    sh.set("dist", str(int(dist_pt * 12700)))
    sh.set("dir", "2700000")
    sh.set("algn", "tl")
    sh.set("rotWithShape", "0")
    c = etree.SubElement(sh, qn("a:srgbClr"))
    c.set("val", color)


def box(sl, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE, shadow=None, dash=None, radius=None):
    s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    _strip_style(s)
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    set_line(s, line, lw, dash)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if shadow:
        hard_shadow(s, shadow)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return s


# ---------------------------------------------------------------- rich text
_TOKEN = re.compile(r"(\*\*.+?\*\*|`.+?`|\{[0-9A-Fa-f]{6}\|.+?\}|\[\[.+?\|.+?\]\])")


def _runs(p, text, font, size, color, bold, italic=False):
    for part in _TOKEN.split(text):
        if not part:
            continue
        f, c, b = font, color, bold
        link = None
        if part.startswith("[[") and part.endswith("]]"):
            part, link = part[2:-2].split("|", 1)
            c = BRAND_INK
        elif part.startswith("**") and part.endswith("**"):
            part, b = part[2:-2], True
            if font == BODY:
                f, b = BODY_SB, False
            elif font == BODY_SB:
                f = BODY  # Segoe UI + bold = the real Segoe UI Bold face
        elif part.startswith("`") and part.endswith("`"):
            part, f = part[1:-1], MONO
        elif part.startswith("{") and part.endswith("}") and "|" in part:
            c, part = part[1:7], part[8:-1]
            if font == BODY:
                f = BODY_SB
        r = p.add_run()
        r.text = part
        r.font.name = f
        r.font.size = Pt(size)
        r.font.bold = b
        r.font.italic = italic
        r.font.color.rgb = rgb(c)
        if link:
            r.hyperlink.address = link if link.startswith("http") else "https://" + link
            r.font.color.rgb = rgb(c)  # hyperlink resets colour to the theme's; keep ours
        # east-asian / complex fallbacks keep ₹ and · in the same face
        rPr = r._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            el = rPr.find(qn(tag))
            if el is None:
                el = etree.SubElement(rPr, qn(tag))
            el.set("typeface", f)


def _bullet(p, color, indent_in=0.16, char="•"):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent_in))))
    pPr.set("indent", str(int(-Inches(indent_in))))
    for tag in ("a:buClr", "a:buSzPct", "a:buFont", "a:buChar", "a:buNone"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bc = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(bc, qn("a:srgbClr")).set("val", color)
    etree.SubElement(pPr, qn("a:buSzPct")).set("val", "100000")
    bf = etree.SubElement(pPr, qn("a:buFont"))
    bf.set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", char)


def fill_text(tf, paras, font=BODY, size=11, color=INK, bold=False, align="l", anchor="t", spacing=1.0,
              after=0.0, bullet=None, italic=False):
    """paras: str | list[str | dict(text=..., size, color, font, bold, bullet, after, align)]"""
    if isinstance(paras, (str, dict)):
        paras = [paras]
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    first = True
    for para in paras:
        spec = para if isinstance(para, dict) else {"text": para}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[spec.get("align", align)]
        p.line_spacing = spec.get("spacing", spacing)
        p.space_after = Pt(spec.get("after", after))
        p.space_before = Pt(spec.get("before", 0))
        b = spec.get("bullet", bullet)
        if b:
            _bullet(p, b if isinstance(b, str) else INK3, spec.get("indent", 0.15))
        _runs(p, spec["text"], spec.get("font", font), spec.get("size", size), spec.get("color", color),
              spec.get("bold", bold), spec.get("italic", italic))
    return tf


def text(sl, x, y, w, h, paras, margin=0.0, **kw):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0)
    fill_text(tf, paras, **kw)
    return tb


def label_box(sl, x, y, w, h, paras, fill=None, line=None, lw=0.75, pad=(0.08, 0.05), shadow=None,
              shape=MSO_SHAPE.RECTANGLE, dash=None, **kw):
    s = box(sl, x, y, w, h, fill=fill, line=line, lw=lw, shape=shape, shadow=shadow, dash=dash)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(pad[0])
    tf.margin_top = tf.margin_bottom = Inches(pad[1])
    fill_text(tf, paras, **kw)
    return s


# ---------------------------------------------------------------- lines & arrows
def _connector_style(c, color, width, dash, head, tail):
    _strip_style(c)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        pd = etree.SubElement(ln, qn("a:prstDash"))
        pd.set("val", dash)
    for tag, kind in (("a:headEnd", head), ("a:tailEnd", tail)):
        if kind:
            e = etree.SubElement(ln, qn(tag))
            e.set("type", kind)
            e.set("w", "med")
            e.set("len", "med")


def line(sl, x1, y1, x2, y2, color=INK, width=1.0, dash=None, head=None, tail=None):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    _connector_style(c, color, width, dash, head, tail)
    return c


def arrow(sl, x1, y1, x2, y2, color=INK, width=1.25, dash=None):
    return line(sl, x1, y1, x2, y2, color, width, dash, tail="triangle")


def polyline(sl, pts, color=INK, width=1.25, dash=None, tail="triangle", head=None):
    """Open freeform path through points (inches) with an arrowhead at the end."""
    fb = sl.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]), scale=1.0)
    fb.add_line_segments([(Inches(px), Inches(py)) for px, py in pts[1:]], close=False)
    s = fb.convert_to_shape()
    _strip_style(s)
    s.fill.background()
    s.line.color.rgb = rgb(color)
    s.line.width = Pt(width)
    ln = _ln(s)
    if dash:
        etree.SubElement(ln, qn("a:prstDash")).set("val", dash)
    for tag, kind in (("a:headEnd", head), ("a:tailEnd", tail)):
        if kind:
            e = etree.SubElement(ln, qn(tag))
            e.set("type", kind)
            e.set("w", "med")
            e.set("len", "med")
    return s


# ---------------------------------------------------------------- pictures & icons
def picture(sl, path, x, y, w=None, h=None, crop=None, line_color=None, lw=0.75, shadow=None):
    kw = {}
    if w is not None:
        kw["width"] = Inches(w)
    if h is not None:
        kw["height"] = Inches(h)
    pic = sl.shapes.add_picture(path, Inches(x), Inches(y), **kw)
    if crop:  # (left, top, right, bottom) fractions
        pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = crop
        if w is not None and h is not None:
            pic.width, pic.height = Inches(w), Inches(h)
    if line_color:
        pic.line.color.rgb = rgb(line_color)
        pic.line.width = Pt(lw)
    if shadow:
        hard_shadow(pic, shadow)
    return pic


def icon_path(comp: str, color: str, lib: str = "lu", stroke: float | None = None) -> str:
    name = f"{lib}-{comp}-{color}{'-s' + str(stroke) if stroke else ''}.png"
    path = os.path.join(ICON_DIR, name)
    if not os.path.exists(path):
        _icon_queue.append({"file": path, "lib": lib, "comp": comp, "color": color, "size": 256, "stroke": stroke})
        render_icons()  # cached on disk, so only the first build pays for this
    return path


def render_icons() -> None:
    if not _icon_queue:
        return
    os.makedirs(ICON_DIR, exist_ok=True)
    spec = os.path.join(ICON_DIR, "_spec.json")
    with open(spec, "w", encoding="utf8") as f:
        json.dump(_icon_queue, f)
    subprocess.run(["node", NODE_ICONS, spec], check=True)
    _icon_queue.clear()


def qr_png(url: str) -> str:
    """QR code (ink on white) for a link, cached next to the icons."""
    import hashlib

    import qrcode

    path = os.path.join(ICON_DIR, "qr-" + hashlib.sha1(url.encode()).hexdigest()[:12] + ".png")
    if not os.path.exists(path):
        os.makedirs(ICON_DIR, exist_ok=True)
        q = qrcode.QRCode(border=0, box_size=20, error_correction=qrcode.constants.ERROR_CORRECT_M)
        q.add_data(url)
        q.make(fit=True)
        q.make_image(fill_color="#" + INK, back_color="white").save(path)
    return path


def icon(sl, comp, x, y, size, color=INK, lib="lu", stroke=None):
    return sl.shapes.add_picture(icon_path(comp, color, lib, stroke), Inches(x), Inches(y), Inches(size), Inches(size))


def icon_tile(sl, comp, x, y, size, fg=WHITE, bg=INK, pad=0.22, lib="lu"):
    """Square tile with a centred icon (square, not a circle — house style)."""
    box(sl, x, y, size, size, fill=bg)
    i = size * pad
    return icon(sl, comp, x + i, y + i, size - 2 * i, fg, lib)


# ---------------------------------------------------------------- composite bits
def pointer(sl, x, y, n, label, w=6.0, size=15, color=INK, chip=INK):
    """Template pointer as a numbered section header: ink square with the number + the pointer text."""
    s = 0.27
    label_box(sl, x, y + 0.01, s, s, {"text": str(n), "align": "c"}, fill=chip, pad=(0, 0), font=DISPLAY, size=12.5,
              color=WHITE, anchor="m")
    return text(sl, x + s + 0.1, y - 0.015, w - s - 0.1, 0.32, label, font=DISPLAY, size=size, color=color, anchor="m")


def chip(sl, x, y, w, h, label, fill=INK, color=WHITE, size=9.5, font=BODY_SB, line=None, align="c"):
    return label_box(sl, x, y, w, h, {"text": label, "align": align}, fill=fill, line=line, pad=(0.05, 0), font=font,
                     size=size, color=color, anchor="m")


def bar_chart(sl, x, y, w, h, cats, values, colors, number_format='0" h"', font_size=10, max_v=None, gap=45):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("NPT", values)
    gf = sl.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.name = BODY
    ch.font.size = Pt(font_size)
    ch.font.color.rgb = rgb(INK2)
    plot = ch.plots[0]
    plot.gap_width = gap
    plot.vary_by_categories = False
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = number_format
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(font_size)
    dl.font.name = BODY_SB
    dl.font.color.rgb = rgb(INK)
    ser = plot.series[0]
    for i, c in enumerate(colors):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(c)
        pt.format.line.fill.background()
    ca = ch.category_axis
    ca.reverse_order = True
    ca.format.line.color.rgb = rgb(RULE)
    ca.has_major_gridlines = False
    ca.tick_labels.font.size = Pt(font_size)
    ca.tick_labels.font.color.rgb = rgb(INK2)
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    if max_v:
        va.maximum_scale = max_v
    va.minimum_scale = 0
    return gf


def set_title(slide, txt, size=None, font=None):
    t = slide.shapes.title
    p = t.text_frame.paragraphs[0]
    runs = p.runs
    runs[0].text = txt
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    if size:
        runs[0].font.size = Pt(size)
    if font:
        runs[0].font.name = font


def set_oval(slide, team_name, color=None):
    for sh in slide.shapes:
        if sh.name.startswith("Oval"):
            tf = sh.text_frame
            ps = tf.paragraphs
            r0 = ps[0].runs[0]
            r0.text = team_name
            for r in ps[0].runs[1:]:
                r._r.getparent().remove(r._r)
            for p in ps[1:]:
                p._p.getparent().remove(p._p)
            if color:
                r0.font.color.rgb = rgb(color)
            return sh


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def find(slide, prefix):
    return [s for s in slide.shapes if s.name.startswith(prefix)]
