"""NWIS — SIH 2026 idea presentation, built on the official SIH2026-IDEA-Presentation-Format.pptx.

    python build.py            -> docs/NWIS_SIH2026_Idea_Presentation.pptx
    python build.py --pdf      -> + docs/NWIS_SIH2026_Idea_Presentation.pdf (exported by the installed PowerPoint)

Fill-in fields (team, theme, links) live in CONFIG below or as flags; a filled-in prototype/video URL gets a QR code.
The hero graphics (cross-section, laptop mockup, architecture) are HTML/SVG rendered by diagrams.mjs (Playwright).
Needs: pip install python-pptx lxml "qrcode[pil]" pywin32; npm install in tools/ (playwright, react-icons, sharp);
PowerPoint for --pdf.
"""
from __future__ import annotations

import os
import subprocess
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from lib import *  # noqa: F401,F403
import lib

ROOT = os.path.dirname(os.path.dirname(lib.HERE))  # repo root (tools/sih_deck/ → ../..)
TEMPLATE = os.path.join(ROOT, "docs", "SIH2026-IDEA-Presentation-Format.pptx")
ARGS = [a for a in sys.argv[1:] if not a.startswith("--")]
OUT = ARGS[0] if ARGS else os.path.join(ROOT, "docs", "NWIS_SIH2026_Idea_Presentation.pptx")
ASSET = lambda n: os.path.join(lib.HERE, "assets", n)  # noqa: E731
GEN = lambda n: os.path.join(lib.HERE, ".render", "gen", n)  # noqa: E731

PLACEHOLDER = "C00000"  # template's own red — makes fill-in fields impossible to miss
CONFIG = {
    "ps_id": "SIH26121",
    "ps_title": "Nearby Wells Intelligence System (NWIS)",
    "theme": None,  # exactly as the SIH portal lists it for SIH26121
    "category": "Software",
    "team_id": None,
    "team_name": None,
    "prototype_url": None,  # deployed site
    "video_url": None,  # YouTube / Drive link of the demo film
    "repo_url": "github.com/aryabailur/sih26121",
}
FLAGS = {"--team": "team_name", "--team-id": "team_id", "--theme": "theme", "--title": "ps_title",
         "--video": "video_url", "--prototype": "prototype_url"}
for a in sys.argv[1:]:
    k, _, v = a.partition("=")
    if k in FLAGS and v:
        CONFIG[FLAGS[k]] = v


def val(key, hint):
    v = CONFIG[key]
    return (v, None) if v else (hint, PLACEHOLDER)


# ---------------------------------------------------------------- deck palette (the app's status/identity colours)
NAVY, NAVY2 = "0E1530", "1A2244"
AMB = "F5A70F"
CYAN, MAG, RED, CRIT, GREEN, GREEN_D = "0EA5E9", "D946EF", "E5383B", "D92D32", "12A679", "0B7A58"
SUB = "2A2F3C"  # secondary text on white — still dark (legibility floor)
ON_DARK = "DCE0EC"
TILE = "F1F3F8"


def t(sl, x, y, w, h, paras, **kw):
    """Body text: Segoe UI Semibold, near-black, unless told otherwise."""
    kw.setdefault("font", BODY_SB)
    kw.setdefault("color", INK)
    return text(sl, x, y, w, h, paras, **kw)


def ptr(sl, x, y, n, label, w=7.0, size=17):
    """A template pointer as a numbered heading: indigo number tile + the pointer text, word for word."""
    s = 0.31
    label_box(sl, x, y, s, s, {"text": str(n), "align": "c"}, fill=BRAND, pad=(0, 0), font=DISPLAY, size=14,
              color=WHITE, anchor="m")
    return text(sl, x + s + 0.12, y - 0.02, w - s - 0.12, s + 0.04, label, font=DISPLAY, size=size, color=INK,
                anchor="m")


def link_card(sl, x, y, w, h, icon_name, label, key, hint):
    v, vc = val(key, hint) if key else (CONFIG["repo_url"], None)
    box(sl, x, y, w, h, fill=WHITE, line=INK, lw=1.25)
    q = h - 0.2
    if vc:
        icon_tile(sl, icon_name, x + 0.1, y + 0.1, q, fg=WHITE, bg=INK, pad=0.2)
        shown = v
    else:
        url = v if v.startswith("http") else "https://" + v
        picture(sl, lib.qr_png(url), x + 0.1, y + 0.1, q, q)
        short = v.rstrip("/").split("/")[-1] if "github.com" in v else v.split("://", 1)[-1]
        shown = f"[[{short[:14] + ('…' if len(short) > 14 else '')} ↗|{url}]]"
    tx = x + q + 0.18
    t(sl, tx, y + 0.08, x + w - tx - 0.04, 0.26, label, font=BODY, bold=True, size=10.5)
    t(sl, tx, y + 0.34, x + w - tx - 0.04, 0.26, shown, size=9.5, color=vc or BRAND_INK)


prs = Presentation(TEMPLATE)

# ---------------------------------------------------------------- drop the "Important instructions" slide
sldIdLst = prs.slides._sldIdLst
last = sldIdLst[-1]
prs.part.drop_rel(last.get(qn("r:id")))
sldIdLst.remove(last)

S = list(prs.slides)
team, team_col = val("team_name", "Team Name")
for s in S[1:]:
    set_oval(s, team, team_col)


# ======================================================================== 1 · TITLE PAGE
def slide_title(sl):
    from pptx.enum.text import PP_ALIGN

    tb = find(sl, "TextBox 9")[0]
    ps = tb.text_frame.paragraphs
    theme, theme_c = val("theme", "‹as listed on the SIH portal›")
    tid, tid_c = val("team_id", "‹Team ID›")
    tname, tname_c = val("team_name", "‹Team name›")
    rows = [
        ("Problem Statement ID – ", CONFIG["ps_id"], None),
        ("Problem Statement Title – ", CONFIG["ps_title"], None),
        ("Theme – ", theme, theme_c),
        ("PS Category – ", CONFIG["category"], None),
        ("Team ID – ", tid, tid_c),
        ("Team Name – ", tname, tname_c),
    ]
    tb.top = Inches(1.95)
    for p, (label, value, colr) in zip(ps[1:], rows):
        r0 = p.runs[0]
        r0.text = label
        r0.font.size = Pt(22)
        r = p.add_run()
        r.text = value
        r.font.size = Pt(21)
        r.font.bold = False
        r.font.name = "Arial"
        if colr:
            r.font.color.rgb = rgb(colr)
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = 1.7
    ps[2].line_spacing = 1.12
    ps[2].space_before = Pt(12)
    ps[2].space_after = Pt(14)
    picture(sl, ASSET("nwis-logo.png"), 7.62, 6.8, 0.46, 0.46)
    text(sl, 8.18, 6.77, 4.6, 0.28, "NWIS · Nearby Wells Intelligence System", font=DISPLAY, size=15, color=INK)
    t(sl, 8.18, 7.05, 4.9, 0.24, "Drilling decision support beside eRTMAC · Oil India Limited", size=10.5, color=SUB)


# ======================================================================== 2 · IDEA
def slide_idea(sl):
    set_title(sl, "NWIS — Nearby Wells Intelligence System", size=30)
    hdr = find(sl, "TextBox 8")[0]
    for p in hdr.text_frame.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    hdr.text_frame.paragraphs[0].runs[0].font.size = Pt(19)
    hdr.left, hdr.top, hdr.width = Inches(0.22), Inches(1.06), Inches(8.3)
    hdr.text_frame._txBody.find(qn("a:bodyPr")).set("wrap", "none")

    L = 0.40
    HW, HH = 748 / 96, 252 / 96  # hero drawn 1:1 (96 px = 1 in) so its type stays at 9–10 pt
    # ---- 1 · the idea, drawn
    ptr(sl, L, 1.56, 1, "Detailed explanation of the proposed solution", w=HW)
    picture(sl, GEN("hero.png"), L, 1.95, HW, HH)
    RX = L + HW + 0.25
    RW = 12.95 - RX
    t(sl, RX, 1.5, RW, 0.44,
      "Links the active well's **bit depth** to what **offset wells hit at that depth** — and warns before the bit "
      "gets there.", size=11.5, spacing=0.98)
    lh = RW * 1304 / 2240
    picture(sl, GEN("laptop.png"), RX, 1.95 + HH - lh, RW, lh)
    chip(sl, RX + 0.16, 1.95 + HH - lh + 0.16, 1.5, 0.26, "LIVE PROTOTYPE", fill=AMB, color=INK, size=9.5, font=BODY_B)

    # ---- 2 · how it addresses the problem
    y2 = 1.95 + HH + 0.16
    ptr(sl, L, y2, 2, "How it addresses the problem", w=5.2)
    chip(sl, L + HW - 3.02, y2 + 0.02, 3.02, 0.28, "✓ All 20 official requirements PASS", fill=GREEN_D, size=10.5,
         font=BODY_B)
    cards = [
        ("LuMapPinned", CYAN, "Nearby wells on a 3D map, any radius"),
        ("LuFileSearch", BRAND, "Offset history in seconds — every line cited"),
        ("LuLayers", MAG, "Correlates depth, formation, casing & mud"),
        ("LuBellRing", CRIT, "Warns before the bit reaches a risky depth"),
    ]
    cg = 0.12
    cw = (HW - cg * 3) / 4
    cy, ch = y2 + 0.4, 0.66
    for k, (ic, col, tx) in enumerate(cards):
        x = L + k * (cw + cg)
        box(sl, x, cy, cw, ch, fill=TILE)
        icon_tile(sl, ic, x + 0.09, cy + 0.12, 0.42, fg=WHITE, bg=col, pad=0.2)
        t(sl, x + 0.6, cy, cw - 0.66, ch, tx, size=10.8, anchor="m", spacing=0.95)
    # links next to the requirement cards
    text(sl, RX, y2 - 0.01, RW, 0.33, "See it working", font=DISPLAY, size=17, color=INK, anchor="m")
    lw3 = [1.42, 1.42]
    lw3.append(RW - sum(lw3) - 0.2)
    x = RX
    for (ic, lab, key, hint), w in zip([("LuMonitor", "Prototype", "prototype_url", "‹add link›"),
                                        ("LuCirclePlay", "Video", "video_url", "‹add link›"),
                                        ("LuGithub", "GitHub", None, None)], lw3):
        link_card(sl, x, cy, w, ch, ic, lab, key, hint)
        x += w + 0.1

    # ---- 3 · innovation: one navy band, the pointer sits inside it
    by = cy + ch + 0.14
    bh = 6.87 - by
    box(sl, L, by, 12.55, bh, fill=NAVY)
    s_ = 0.31
    label_box(sl, L + 0.16, by + 0.14, s_, s_, {"text": "3", "align": "c"}, fill=AMB, pad=(0, 0), font=DISPLAY,
              size=14, color=INK, anchor="m")
    text(sl, L + 0.56, by + 0.1, 1.95, bh - 0.2, "Innovation and uniqueness of the solution", font=DISPLAY, size=15,
         color=WHITE, anchor="m", spacing=0.95)
    usp = [
        ("LuCrosshair", "Depth-linked", "bit depth drives map, 3D, risk, search"),
        ("LuQuote", "Evidence-first", "every answer cites report, page, depth"),
        ("LuScale", "Explainable + ML", "6-factor score + ML check (AUC 0.86)"),
        ("LuClipboardList", "Look-ahead brief", "next 300 m for the shift handover, ₹ NPT"),
        ("LuBadgeCheck", "Real-data proof", "1,970 real well histories, 94 % precision"),
    ]
    ux0 = L + 2.6
    uw = (L + 12.55 - ux0) / 5
    for k, (ic, ti, sub) in enumerate(usp):
        x = ux0 + k * uw
        line(sl, x, by + 0.14, x, by + bh - 0.14, "3A4570", 1.0)
        icon(sl, ic, x + 0.14, by + 0.15, 0.28, AMB)
        t(sl, x + 0.48, by + 0.12, uw - 0.5, 0.34, ti, font=BODY, bold=True, size=11.8, color=WHITE, anchor="m")
        t(sl, x + 0.16, by + 0.5, uw - 0.24, bh - 0.52, sub, size=10.2, color=ON_DARK, spacing=0.93)


# ======================================================================== 3 · TECHNICAL APPROACH
def slide_tech(sl):
    for s_ in find(sl, "TextBox 8"):
        remove_shape(s_)
    L = 0.40
    ptr(sl, L, 1.3, 1, "Technologies to be used", w=5)
    groups = [
        ("Frontend", CYAN, [("si", "SiNextdotjs", "000000", "Next.js 16"), ("si", "SiReact", "149ECA", "React 19"),
                            ("si", "SiMaplibre", "396CB2", "MapLibre"), ("si", "SiThreedotjs", "000000", "three.js")]),
        ("Backend", GREEN_D, [("si", "SiPython", "3776AB", "Python 3"), ("si", "SiFastapi", "009688", "FastAPI"),
                              ("si", "SiPydantic", "E92063", "Pydantic")]),
        ("AI · NLP · OCR", BRAND, [("lu", "LuScanText", "5B4BFF", "Tesseract"), ("lu", "LuBrainCircuit", "C026D3", "Drilling NLP"),
                                   ("lu", "LuSearch", "0284C7", "BM25 + vec"), ("lu", "LuSigma", "0B7A58", "Log. regr."),
                                   ("si", "SiAnthropic", "191919", "Claude opt.")]),
        ("Data & integration", "B45309", [("si", "SiPostgresql", "4169E1", "PostgreSQL"), ("si", "SiSqlite", "003B57", "SQLite"),
                                          ("lu", "LuRadio", "B45309", "WITSML")]),
    ]
    n = sum(len(g[2]) for g in groups)
    gg, ig = 0.2, 0.06
    iw = (12.55 - gg * (len(groups) - 1) - ig * (n - len(groups))) / n
    x = L
    gy, iy, ih = 1.64, 1.84, 0.52
    for gname, gcol, items in groups:
        gw = len(items) * iw + (len(items) - 1) * ig
        t(sl, x, gy, gw, 0.2, gname, font=BODY, bold=True, size=10, color=gcol)
        for libn, comp, col, name in items:
            box(sl, x, iy, iw, ih, fill=WHITE, line="C9CEDA", lw=0.75)
            icon(sl, comp, x + (iw - 0.24) / 2, iy + 0.05, 0.24, col, libn)
            t(sl, x + 0.01, iy + 0.3, iw - 0.02, 0.2, {"text": name, "align": "c"}, size=9)
            x += iw + ig
        x += gg - ig

    ptr(sl, L, 2.46, 2, "Methodology and process for implementation", w=8)
    aw = 12.25
    picture(sl, GEN("arch.png"), L + (12.55 - aw) / 2, 2.84, aw, aw * 392 / 1205)


# ======================================================================== 4 · FEASIBILITY
def slide_feasibility(sl):
    for s_ in find(sl, "TextBox 8"):
        remove_shape(s_)
    L, LW = 0.40, 4.1
    ptr(sl, L, 1.3, 1, "Analysis of the feasibility of the idea", w=LW + 0.4)
    stats = [
        ("20/20", "official SIH26121 requirements pass an API check", GREEN_D, WHITE),
        ("94 %", "precision on real, held-out well histories", BRAND, WHITE),
        ("0.86", "AUC of the learned risk cross-check", NAVY, WHITE),
        ("< 30 ms", "for a cited answer or a full risk evaluation", AMB, INK),
    ]
    sw, sh, sg = (LW - 0.14) / 2, 1.14, 0.14
    for k, (num, lab, bg, fg) in enumerate(stats):
        x = L + (k % 2) * (sw + sg)
        y = 1.74 + (k // 2) * (sh + sg)
        box(sl, x, y, sw, sh, fill=bg)
        text(sl, x + 0.14, y + 0.08, sw - 0.2, 0.52, num, font=DISPLAY, size=30, color=fg)
        t(sl, x + 0.14, y + 0.62, sw - 0.22, 0.54, lab, size=10.5, color=fg, spacing=0.95)
    checks = [("Technical", "open source, one server, no GPU"),
              ("Operational", "beside eRTMAC, no workflow change"),
              ("Economic", "no licence cost; saves NPT days")]
    for k, (a, b) in enumerate(checks):
        y = 4.36 + k * 0.34
        icon(sl, "LuCircleCheck", L, y + 0.04, 0.26, GREEN_D)
        t(sl, L + 0.36, y, LW - 0.36, 0.34, f"**{a}** — {b}", size=11, anchor="m")

    # ---- challenges → strategies
    R = 4.8
    cw1, sx = 3.3, 8.52
    sw2 = 12.95 - sx
    ptr(sl, R, 1.3, 2, "Potential challenges and risks", w=cw1 + 0.4)
    ptr(sl, sx, 1.3, 3, "Strategies for overcoming these challenges", w=sw2)
    rows = [
        ("LuLock", "Confidential well data", "High", "On-prem in OIL's network, **runs offline**; the LLM is optional and **off by default**; SSO roles."),
        ("LuFileWarning", "Scanned, messy old reports", "High", "OCR with per-page confidence; **an engineer approves** every extracted event."),
        ("LuBellRing", "Alert fatigue on the rig", "Medium", "Weak windows need **live confirmation**; one alert per window — it **escalates, never repeats**."),
        ("LuScale", "Trust in AI advice", "Medium", "Every alert shows **Why?** — factors, cited evidence, confidence. **The engineer decides.**"),
        ("LuSigma", "Demo-tuned weights", "Low", "**Refit on OIL's NPT history**, then a **shadow pilot** before alerts go live."),
    ]
    lvl = {"High": CRIT, "Medium": "C2410C", "Low": GREEN_D}
    ry, rh, rg = 1.74, 0.62, 0.1
    for k, (ic, ch, lv, st) in enumerate(rows):
        y = ry + k * (rh + rg)
        box(sl, R, y, cw1, rh, fill="FCEBEB")
        icon_tile(sl, ic, R + 0.09, y + 0.1, rh - 0.2, fg=WHITE, bg=lvl[lv], pad=0.2)
        t(sl, R + rh + 0.05, y, cw1 - rh - 0.85, rh, ch, font=BODY, bold=True, size=11.5, anchor="m", spacing=0.95)
        chip(sl, R + cw1 - 0.78, y + (rh - 0.26) / 2, 0.68, 0.26, lv, fill=lvl[lv], size=9.5, font=BODY_B)
        arrow(sl, R + cw1 + 0.05, y + rh / 2, sx - 0.06, y + rh / 2, INK, 2.0)
        box(sl, sx, y, sw2, rh, fill="E4F5EE")
        icon(sl, "LuCircleCheck", sx + 0.12, y + (rh - 0.26) / 2, 0.26, GREEN_D)
        t(sl, sx + 0.48, y, sw2 - 0.58, rh, st, size=10.8, anchor="m", spacing=0.95)

    # ---- viability roadmap
    vy = 5.44
    text(sl, L, vy, 9, 0.32, "Viability — from prototype to OIL rigs in 12 months", font=DISPLAY, size=17, color=INK,
         anchor="m")
    phases = [
        ("NOW", "Prototype ready", "20/20 requirements · real-data proof · demo film", GREEN_D, WHITE),
        ("MONTH 0–3", "Ingest one field", "OIL DDR / WCR archive, SME review, refit weights", BRAND, WHITE),
        ("MONTH 3–6", "Shadow pilot at eRTMAC", "2 rigs · alerts logged, not shown · measure lead", NAVY, WHITE),
        ("MONTH 6–12", "Roll out", "live WITSML feed · SSO · all Assam rigs", AMB, INK),
    ]
    pw = 12.55 / 4
    py, ph = vy + 0.4, 6.86 - (vy + 0.4)
    for k, (when, ti, d, bg, fg) in enumerate(phases):
        x = L + k * pw
        shp = box(sl, x, py, pw + (0.18 if k < 3 else 0), ph, fill=bg,
                  shape=MSO_SHAPE.PENTAGON if k == 0 else MSO_SHAPE.CHEVRON)
        shp.adjustments[0] = 0.28
        tx = x + (0.2 if k == 0 else 0.46)
        t(sl, tx, py + 0.07, pw - 0.62, 0.22, when, font=BODY, bold=True, size=9.5, color=fg)
        text(sl, tx, py + 0.27, pw - 0.62, 0.3, ti, font=DISPLAY, size=15, color=fg)
        t(sl, tx, py + 0.58, pw - 0.62, 0.42, d, size=9.6, color=fg, spacing=0.93)
    # draw order: later chevrons must not cover earlier text — shapes were added in order, text after each, fine


# ======================================================================== 5 · IMPACT
def slide_impact(sl):
    for s_ in find(sl, "TextBox 8"):
        remove_shape(s_)
    L, LW = 0.40, 4.15
    ptr(sl, L, 1.3, 1, "Potential impact on the target audience", w=LW + 0.4)
    people = [
        ("LuHardHat", BRAND, "Rig engineer & driller", "Sees what offsets hit at this depth — 50 m before the bit gets there."),
        ("LuMonitor", CYAN, "eRTMAC monitoring team", "Live data linked to offset history; cited evidence in seconds."),
        ("LuClipboardList", "C2410C", "Planners & DWOP", "Look-ahead brief: hazards, mud window and NPT in ₹ per section."),
        ("LuGraduationCap", GREEN_D, "New engineers & managers", "The field's memory stays searchable; every alert is audited."),
    ]
    for k, (ic, col, who, what) in enumerate(people):
        y = 1.76 + k * 1.02
        icon_tile(sl, ic, L, y + 0.04, 0.64, fg=WHITE, bg=col, pad=0.2)
        text(sl, L + 0.8, y - 0.02, LW - 0.8, 0.32, who, font=DISPLAY, size=15.5, color=INK)
        t(sl, L + 0.8, y + 0.31, LW - 0.8, 0.56, what, size=11, spacing=0.97)
    label_box(sl, L, 5.92, LW, 0.92,
              "Engineers don't lack data. They lack the link between **this depth, right now** and **what happened "
              "here before**.", fill=BRAND_SOFT, pad=(0.14, 0.06), font=BODY_SB, size=11.5, color=INK, anchor="m",
              spacing=0.97)

    # ---- centre: the numbers
    M, MW = 4.85, 3.85
    box(sl, M, 1.3, MW, 5.56, fill=NAVY)
    px, pw = M + 0.25, MW - 0.45
    t(sl, px, 1.42, pw, 0.26, "Why it matters", font=BODY, bold=True, size=11, color=AMB)
    big = [
        ("≈ 20 %", "of well-construction time is lost as non-productive time (NPT) [2]", WHITE),
        ("up to 40 %", "of NPT comes from kicks, losses and wellbore instability — exactly what NWIS warns about [3]", WHITE),
        ("≈ ₹21 cr / yr", "if NWIS saves one NPT day on each of the ~70 wells OIL spuds a year [4]*", AMB),
    ]
    y = 1.72
    for k, (n, d, c) in enumerate(big):
        if k:
            line(sl, px, y - 0.1, px + pw, y - 0.1, "3A4570", 1.0)
        text(sl, px, y, pw, 0.7, n, font=DISPLAY, size=40, color=c)
        t(sl, px, y + 0.72, pw, 0.7, d, size=11, color=WHITE, spacing=0.97)
        y += 1.6
    t(sl, px, 6.42, pw, 0.4, "* 1 day × 70 wells × ₹30 lakh/day assumed spread rate — an illustration, not a forecast.",
      size=8.8, color=ON_DARK, spacing=0.93)

    # ---- right: benefits
    R, RW = 8.95, 4.0
    ptr(sl, R, 1.3, 2, "Benefits of the solution", w=RW)
    ben = [
        ("LuIndianRupee", "C2410C", "Economic", "Fewer repeated losses, stuck pipe and kicks → less NPT, faster planning."),
        ("LuShieldCheck", CRIT, "Safety", "Earlier kick and overpressure awareness supports well control."),
        ("LuLeaf", GREEN_D, "Environmental", "Fewer loss events → less mud, LCM and cement lost; less rig diesel."),
        ("LuUsers", BRAND, "Social", "Expert knowledge stays in India's energy PSU; faster onboarding."),
    ]
    bh, bg = 1.17, 0.12
    for k, (ic, col, ti, d) in enumerate(ben):
        y = 1.74 + k * (bh + bg)
        box(sl, R, y, RW, bh, fill=TILE)
        icon_tile(sl, ic, R + 0.12, y + 0.14, 0.5, fg=WHITE, bg=col, pad=0.2)
        text(sl, R + 0.76, y + 0.12, RW - 0.86, 0.34, ti, font=DISPLAY, size=16, color=col)
        t(sl, R + 0.76, y + 0.46, RW - 0.86, 0.68, d, size=10.8, spacing=0.96)


# ======================================================================== 6 · RESEARCH
def slide_research(sl):
    for s_ in find(sl, "TextBox 8"):
        remove_shape(s_)
    set_title(sl, "RESEARCH AND REFERENCES")
    L = 0.40
    ptr(sl, L, 1.3, 1, "Research — where today's approaches fall short", w=7)
    cols = [["Manual", "offset review"], ["Monitoring", "alone"], ["Generic", "chatbot"], ["NWIS"]]
    feats = [
        ("Tied to the live bit depth", [0, 2, 0, 2]),
        ("Reads DDR / WCR text, even scans", [2, 0, 1, 2]),
        ("Cites report page + depth", [1, 0, 0, 2]),
        ("Warns before the risky depth", [0, 1, 0, 2]),
    ]
    fw, cw = 2.9, 0.97
    hy, hh, rh = 1.74, 0.56, 0.5
    for k, c in enumerate(cols):
        x = L + fw + k * cw
        hot = k == 3
        box(sl, x, hy, cw - 0.06, hh, fill=BRAND if hot else "E3E6EE")
        text(sl, x, hy, cw - 0.06, hh, [{"text": s_, "align": "c"} for s_ in c], font=DISPLAY if hot else BODY_SB,
             size=15 if hot else 10, color=WHITE if hot else INK, anchor="m", spacing=0.9)
    marks = {2: ("LuCircleCheck", GREEN_D), 1: ("LuCircleDot", "C2410C"), 0: ("LuCircleX", "9AA0AE")}
    for r, (f, vals) in enumerate(feats):
        y = hy + hh + 0.06 + r * rh
        if r % 2 == 0:
            box(sl, L, y, fw + 4 * cw - 0.06, rh, fill=TILE)
        t(sl, L + 0.1, y, fw - 0.12, rh, f, size=11.5, anchor="m")
        for k, v in enumerate(vals):
            ic, col = marks[v]
            icon(sl, ic, L + fw + k * cw + (cw - 0.06 - 0.28) / 2, y + (rh - 0.28) / 2, 0.28, col)
    for k, (v, lab) in enumerate([(2, "yes"), (1, "partly"), (0, "no")]):
        ic, col = marks[v]
        icon(sl, ic, L + 0.1 + k * 0.85, hy + 0.2, 0.18, col)
        t(sl, L + 0.33 + k * 0.85, hy + 0.17, 0.6, 0.24, lab, size=10, color=SUB)

    # ---- right: real-data proof
    R, RW = 7.7, 5.25
    ptr(sl, R, 1.3, 2, "Validated on real public well records", w=RW)
    py0, ph = 1.74, 2.8
    box(sl, R, py0, RW, ph, fill=NAVY)
    t(sl, R + 0.22, py0 + 0.14, RW - 0.4, 0.3, "Same extractor, rules frozen — Norwegian shelf [7]",
      font=BODY, bold=True, size=11.5, color=WHITE)
    nums = [("1,970", "well histories read", WHITE), ("611", "problems found", AMB),
            ("94 %", "held-out precision", "4FDCAA")]
    nw = (RW - 0.44) / 3
    for k, (n, d, c) in enumerate(nums):
        x = R + 0.22 + k * nw
        text(sl, x, py0 + 0.56, nw - 0.1, 0.7, n, font=DISPLAY, size=38, color=c)
        t(sl, x, py0 + 1.3, nw - 0.12, 0.5, d, size=10.5, color=WHITE, spacing=0.95)
    fam = [("Stuck pipe 161", 161, MAG), ("Mud loss 144", 144, CYAN), ("Kick 112", 112, RED), ("NPT 100", 100, "8B93A7"),
           ("Instability 82", 82, "14B8A6"), ("", 12, "EAB308")]
    bx, bw = R + 0.22, RW - 0.44
    tot = sum(v for _, v, _ in fam)
    by = py0 + 1.78
    for name, v, c in fam:
        w = bw * v / tot
        box(sl, bx, by, max(w - 0.02, 0.02), 0.12, fill=c)
        bx += w
    t(sl, R + 0.22, by + 0.18, RW - 0.44, 0.3, "stuck pipe 161 · mud loss 144 · kick 112 · NPT 100 · instability 82",
      size=9.6, color=ON_DARK)
    t(sl, R + 0.22, by + 0.52, RW - 0.44, 0.5, "Prior work mines report text [5] or predicts stuck pipe from parameters "
      "[6]. **NWIS joins both at the live bit depth.**", size=10.2, color=WHITE, spacing=0.95)

    # ---- references
    ry = 4.76
    ptr(sl, L, ry, 3, "Details / links of the reference and research work", w=9)
    refs = REFS["short"]
    colw = (12.55 - 0.4) / 2
    per = (len(refs) + 1) // 2
    for c in range(2):
        chunk = refs[c * per:(c + 1) * per]
        t(sl, L + c * (colw + 0.4), ry + 0.44, colw, 1.8, [{"text": r, "after": 8} for r in chunk], size=11,
          spacing=0.95)


NOTES = [
    "SIH26121 from Oil India Limited: a Nearby Wells Intelligence System. One sentence: NWIS gives every drilling "
    "engineer the field's memory at the depth it matters.",
    "The picture is the idea. The active well is at 3,100 m. Offset wells hit mud loss at 3,150, stuck pipe at 3,380 "
    "and a kick at 3,580. NWIS reads their reports, links each event to the same depth on the active well and raises "
    "an alert 50 m before the bit gets there — with the fix that worked and the source page. The laptop is the working "
    "prototype; all 20 official requirements pass an automated check.",
    "Top: the stack, all open source. Below: OIL's reports go through OCR and NLP, an engineer approves what becomes "
    "knowledge, and a depth-indexed knowledge base feeds the intelligence engine behind one API. The amber line is the "
    "live loop — every eRTMAC depth step re-scores the windows ahead, and the chevrons show how an alert is made.",
    "Feasibility is shown, not claimed: it is built and measured. Each risk has a concrete answer, most already built. "
    "The path to the rig is one field ingested, then a shadow pilot at eRTMAC before any alert goes live.",
    "Industry numbers first: about a fifth of well time is NPT and up to 40 % of it comes from kicks, losses and "
    "instability. At OIL's scale one avoided NPT day per well is about 21 crore rupees a year — an illustration at an "
    "assumed spread rate, not a forecast.",
    "NWIS is the only approach tied to the live bit depth with cited evidence. We proved the extractor on 1,970 real "
    "public well histories with 94 % precision. Give us three anonymised OIL DDR pages and we ingest them live.",
]

REFS = {}

if __name__ == "__main__":
    import refs  # noqa: E402

    REFS.update(refs.REFS)
    if "--no-render" not in sys.argv:
        subprocess.run(["node", os.path.join(lib.HERE, "diagrams.mjs")], check=True, cwd=lib.HERE)
    slide_title(S[0])
    slide_idea(S[1])
    slide_tech(S[2])
    slide_feasibility(S[3])
    slide_impact(S[4])
    slide_research(S[5])
    for s, note in zip(S, NOTES):
        s.notes_slide.notes_text_frame.text = note
    lib.render_icons()
    prs.save(OUT)
    print("saved", OUT)
    if "--pdf" in sys.argv:
        pdf = os.path.abspath(os.path.splitext(OUT)[0] + ".pdf")
        try:  # print-intent export keeps every hyperlink (SaveAs-PDF drops some) — needs pywin32
            import win32com.client

            app = win32com.client.Dispatch("PowerPoint.Application")
            try:
                deck = app.Presentations.Open(os.path.abspath(OUT), True, False, False)
                # PDF, print intent, no frames, handout order, slides, no hidden, range None, all slides, "",
                # doc properties, IRM, structure tags (accessible PDF), bitmap missing fonts, not PDF/A
                deck.ExportAsFixedFormat(pdf, 2, 2, 0, 1, 1, 0, None, 1, "", True, True, True, True, False)
                deck.Close()
            finally:
                app.Quit()
        except ImportError:
            render = os.path.join(ROOT, "tools", "render_pptx.ps1")
            subprocess.run(["powershell", "-ExecutionPolicy", "Bypass", "-File", render, "-Deck", os.path.abspath(OUT),
                            "-OutDir", os.path.join(lib.HERE, ".render"), "-Pdf", pdf], check=True)
        print("saved", pdf)
